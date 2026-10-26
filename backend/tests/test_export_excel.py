"""Excel customer export service and endpoint verification tests."""

import io
from fastapi.testclient import TestClient
import openpyxl
import pytest

from app.services.calculation import CalculationService
from app.services.export_excel import ExcelExportService


class TestExcelCustomerExport:
    """Verifies that generated customer workbooks are password-free, unlocked, and formula-preserving."""

    @pytest.fixture
    def sample_quote_and_calc(self):
        quote = {
            "enquiry_no": "MKTG/ENQ/26-27/TEST-EXPORT",
            "customer_name": "Antigravity Engineering Ltd",
            "project_name": "Pipeline Water Project",
            "delivery_location": "Bagodar, Jharkhand",
            "trip_rate": 93000.0,
        }
        lines = [
            {
                "size_mm": 1226.0,
                "thickness_mm": 8.4,
                "basis": "ID",
                "quantity": 2000.0,
                "unit": "Metres",
                "steel_rate": 62000.0,
                "conversion": 4000.0,
                "ext_rate": 500.0,
                "int_rate": 250.0,
                "trip_rate": 93000.0,
                "handling": 100.0,
                "financial_pct": 1.0,
                "manual_fc": 1024.0,
            }
        ]
        calc = CalculationService.calculate_quotation(lines, trip_rate=93000.0, mode="REFERENCE")
        return quote, calc

    def test_direct_workbook_generation_properties(self, sample_quote_and_calc):
        """Generated workbook must be password-free, unlocked, and retain local Excel formulas."""
        quote, calc = sample_quote_and_calc

        stream = ExcelExportService.generate_quotation_workbook(quote, calc)
        assert isinstance(stream, io.BytesIO)

        # Open workbook with openpyxl
        wb = openpyxl.load_workbook(stream)

        # 1. Structure check: sheets Offer and Synopsis
        sheet_names = wb.sheetnames
        assert "Offer" in sheet_names
        assert "Synopsis" in sheet_names

        # 2. Protection check: No password, protection disabled
        assert wb.security.lockStructure is False
        for ws in wb.worksheets:
            assert ws.protection.sheet is False

        # 3. All cells unlocked
        ws_offer = wb["Offer"]
        for row in ws_offer.iter_rows(min_row=1, max_row=20, min_col=1, max_col=9):
            for cell in row:
                if cell.value is not None:
                    assert cell.protection.locked is False

        # 4. Preserved local Excel formulas
        # Line 1 amount: row 9 col I
        assert str(ws_offer["I9"].value).startswith("=F9*H9")

        # Subtotal formula: row 10 col I
        assert str(ws_offer["I10"].value).startswith("=SUM(I9:I9)")

        # GST formula: row 11 col I
        assert str(ws_offer["I11"].value).startswith("=ROUND(I10*0.18")

        # Grand Total formula: row 12 col I
        assert str(ws_offer["I12"].value).startswith("=I10+I11")

        # 5. External link check: ensure no external bracket references in cell formulas
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    val = str(cell.value or "")
                    if val.startswith("="):
                        assert "[" not in val and "]" not in val

    def test_export_api_endpoint(self, authenticated_client: TestClient):
        """Export API endpoint returns streamed .xlsx file with proper headers."""
        # Create and calculate quotation
        create_resp = authenticated_client.post(
            "/api/quotations",
            json={
                "enquiry_no": "MKTG/ENQ/26-27/EXP-01",
                "customer_name": "Export Test Customer",
                "lines": [
                    {
                        "line_number": 1,
                        "size_mm": 1226.0,
                        "thickness_mm": 8.4,
                        "basis": "ID",
                        "quantity": 500.0,
                        "steel_rate": 62000.0,
                        "conversion": 4000.0,
                        "ext_rate": 500.0,
                        "int_rate": 250.0,
                        "manual_fc": 1024.0,
                    }
                ],
            },
        )
        quote_id = create_resp.json()["id"]

        # 1. Export uncalculated quotation fails with HTTP 400
        uncalc_export = authenticated_client.get(f"/api/quotations/{quote_id}/export")
        assert uncalc_export.status_code == 400

        # 2. Calculate
        authenticated_client.post(f"/api/quotations/{quote_id}/calculate")

        # 3. Export calculated quotation succeeds with attachment
        export_resp = authenticated_client.get(f"/api/quotations/{quote_id}/export")
        assert export_resp.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in export_resp.headers["content-type"]
        assert "attachment" in export_resp.headers["content-disposition"]

        # Validate streamed content is loadable as an Excel workbook
        content_stream = io.BytesIO(export_resp.content)
        wb = openpyxl.load_workbook(content_stream)
        assert "Offer" in wb.sheetnames
