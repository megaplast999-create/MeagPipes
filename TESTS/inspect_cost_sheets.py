"""
Inspect the 3 cost sheets in Downloads:
1. MKTGENQ26-27594 - Unipro.xlsx
2. MKTGENQ26-27-594 - Unipro 2.xlsx
3. MKTGENQ26-27610 - Lodha 1.xlsx
"""

import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

files = [
    r'C:\Users\Administrator\Downloads\MKTGENQ26-27594 - Unipro.xlsx',
    r'C:\Users\Administrator\Downloads\MKTGENQ26-27-594 - Unipro 2.xlsx',
    r'C:\Users\Administrator\Downloads\MKTGENQ26-27610 - Lodha 1.xlsx'
]

for fpath in files:
    fname = os.path.basename(fpath)
    print(f"\n{'='*80}")
    print(f"FILE: {fname}")
    print(f"{'='*80}")
    
    if not os.path.exists(fpath):
        print("  FILE NOT FOUND!")
        continue
        
    wb_val = openpyxl.load_workbook(fpath, data_only=True)
    wb_form = openpyxl.load_workbook(fpath, data_only=False)
    
    for sname in wb_val.sheetnames:
        print(f"\n--- SHEET: {sname} ---")
        ws_v = wb_val[sname]
        ws_f = wb_form[sname]
        
        max_r = ws_v.max_row
        max_c = ws_v.max_column
        print(f"Dimensions: {max_r} rows x {max_c} columns")
        
        for r in range(1, max_r + 1):
            row_vals = [ws_v.cell(r, c).value for c in range(1, max_c + 1)]
            row_forms = [ws_f.cell(r, c).value for c in range(1, max_c + 1)]
            
            # Check if row has any non-empty cells
            if any(v is not None for v in row_vals):
                cells_str = []
                for c, (v, f) in enumerate(zip(row_vals, row_forms), 1):
                    if v is not None or f is not None:
                        col_letter = openpyxl.utils.get_column_letter(c)
                        if str(f).startswith('='):
                            cells_str.append(f"{col_letter}{r}: [{f}] = {v}")
                        else:
                            cells_str.append(f"{col_letter}{r}: {v}")
                print(f"Row {r:2d}: " + " | ".join(cells_str))

print("\nFinished inspection.")
