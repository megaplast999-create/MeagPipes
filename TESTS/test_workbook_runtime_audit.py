from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
import openpyxl


NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "DEV" / "Mega_Pipes_Quotation_DEV.xlsm"
EXPECTED_SHEETS = [
    "COMPARE", "START", "QUOTE", "QUOTE_LINES", "COSTING", "OFFER", "MASTERS",
    "SETTINGS", "CUSTOMERS", "PLANTS", "DISTANCE_CACHE", "QUOTE_LOG",
]
EXPECTED_TABLES = {
    "tblProducts", "tblSizes", "tblThickness", "tblGrades",
    "tblExternalCoatings", "tblInternalCoatings", "tblPlants",
    "tblCustomers", "tblPaymentTerms", "tblSettings", "tblQuoteLines",
    "tblDistanceCache", "tblQuoteLog", "tblPlantProducts",
    "tblStacking", "tblSlabs", "tblCompPrices",
}


def test_macro_enabled_package_and_structure():
    with zipfile.ZipFile(WORKBOOK) as archive:
        assert "xl/vbaProject.bin" in archive.namelist()
        content_types = archive.read("[Content_Types].xml").decode("utf-8")
        assert "application/vnd.ms-excel.sheet.macroEnabled.main+xml" in content_types
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        sheets = [sheet.attrib["name"] for sheet in workbook.find("m:sheets", NS)]
        assert sheets == EXPECTED_SHEETS
        table_names = {
            ET.fromstring(archive.read(name)).attrib["name"]
            for name in archive.namelist()
            if name.startswith("xl/tables/table") and name.endswith(".xml")
        }
        assert table_names == EXPECTED_TABLES


def test_customer_master_has_one_nonblank_customer():
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb["CUSTOMERS"]
    tbl = ws.tables["tblCustomers"]
    min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(tbl.ref)
    customer_ids = [
        ws.cell(r, min_col).value
        for r in range(min_row + 1, max_row + 1)
        if ws.cell(r, min_col).value
    ]
    assert customer_ids == ["C001"]
